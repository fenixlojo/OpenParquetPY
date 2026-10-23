from datetime import datetime
from pathlib import Path
import io
import json
import tempfile
import threading
import time
import unittest
import urllib.error
import urllib.request
import duckdb
import openpyxl
import app
import conversion


class ConversionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server=app.ThreadingHTTPServer(('127.0.0.1',0),app.Handler)
        cls.thread=threading.Thread(target=cls.server.serve_forever,daemon=True)
        cls.thread.start()
        cls.base=f'http://127.0.0.1:{cls.server.server_port}'

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown();cls.server.server_close();cls.thread.join()

    def post(self,path,value):
        request=urllib.request.Request(self.base+path,data=json.dumps(value).encode(),headers={'Content-Type':'application/json'})
        with urllib.request.urlopen(request) as response:return json.load(response)

    def source(self,name,body):
        from urllib.parse import quote
        request=urllib.request.Request(self.base+'/api/convert/source',data=body,headers={'X-Filename':quote(name)})
        with urllib.request.urlopen(request) as response:return json.load(response)

    def convert(self,source,**options):
        job=self.post('/api/convert/start',dict(token=source['token'],options=options))['job']
        deadline=time.monotonic()+15
        while time.monotonic()<deadline:
            with urllib.request.urlopen(self.base+'/api/convert/job?id='+job) as response:status=json.load(response)
            if status['state']!='running':return status
            time.sleep(.02)
        self.fail('Conversion timed out')

    def discard(self,source):
        self.post('/api/convert/discard',{'token':source['token']})

    def test_csv_download_registration_and_savings(self):
        body=('persona_id;nombre\n'+''.join(f'{42325+i};persona repetida\n' for i in range(6000))).encode()
        source=self.source('datos.csv',body)
        try:
            status=self.convert(source,header=True)
            self.assertEqual(status['state'],'complete',status)
            result=status['result'];self.assertEqual(result['originalSize'],len(body));self.assertLess(result['outputSize'],len(body))
            with urllib.request.urlopen(self.base+result['downloadUrl']) as response:
                binary=response.read();self.assertIn('attachment',response.headers['Content-Disposition'])
            self.assertEqual(binary[:4],b'PAR1');self.assertEqual(len(binary),result['outputSize'])
            with urllib.request.urlopen(self.base+'/api/rows?id='+result['info']['id']) as response:rows=json.load(response)
            self.assertEqual(rows['total'],6000);self.assertEqual(rows['rows'][0],[42325,'persona repetida'])
            self.post('/api/close',{'id':result['info']['id']})
            with urllib.request.urlopen(self.base+result['downloadUrl']) as response:
                self.assertEqual(response.read(),binary)
        finally:self.discard(source)

    def test_csv_encoding_no_header_and_small_file_increase(self):
        body='00123,José\n00456,María\n'.encode('cp1252');source=self.source('acentos.csv',body)
        try:
            status=self.convert(source,header=False,allText=True,encoding='windows-1252')
            self.assertEqual(status['state'],'complete',status)
            result=status['result'];self.assertGreater(result['outputSize'],result['originalSize'])
            path,_=app.FILES[result['info']['id']]
            with duckdb.connect() as db:rows=db.execute('SELECT * FROM read_parquet(?)',[str(path)]).fetchall()
            self.assertEqual(rows,[('00123','José'),('00456','María')])
        finally:self.discard(source)

    def test_excel_sheet_selection_dates_and_headers(self):
        book=openpyxl.Workbook();book.active.title='Otra';book.active.append(['ignorar'])
        tab=book.create_sheet('Elegida');tab.append(['id','nombre','fecha']);tab.append([42325,'José',datetime(2026,10,23,12,30)])
        data=io.BytesIO();book.save(data);book.close();source=self.source('libro.xlsx',data.getvalue())
        try:
            self.assertEqual(source['sheets'],['Otra','Elegida'])
            status=self.convert(source,header=True,sheet='Elegida');self.assertEqual(status['state'],'complete',status)
            result=status['result'];self.assertEqual(result['name'],'libro_Elegida.parquet')
            path,_=app.FILES[result['info']['id']]
            with duckdb.connect() as db:rows=db.execute('SELECT * FROM read_parquet(?)',[str(path)]).fetchall()
            self.assertEqual(rows,[(42325,'José',datetime(2026,10,23,12,30))])
        finally:self.discard(source)

    def test_empty_formula_and_invalid_sheet_errors(self):
        book=openpyxl.Workbook();book.active.title='Vacía';tab=book.create_sheet('Fórmulas');tab.append(['valor']);tab.append(['=1+1'])
        data=io.BytesIO();book.save(data);book.close();source=self.source('libro.xlsx',data.getvalue())
        try:
            for sheet in ['Vacía','Fórmulas','No existe']:
                with self.subTest(sheet=sheet):self.assertEqual(self.convert(source,header=True,sheet=sheet)['state'],'error')
        finally:self.discard(source)

    def test_unsupported_and_empty_uploads(self):
        for name,body in [('archivo.txt',b'text'),('vacio.csv',b'')]:
            with self.subTest(name=name),self.assertRaises(urllib.error.HTTPError):self.source(name,body)


if __name__=='__main__':unittest.main()
